from datetime import datetime
from pathlib import Path
import os
import sys

# Add root to path for utility access
sys.path.append(os.path.abspath(os.path.join(os.path.dirname(__file__), '../../')))
from utils.notify import notify_user

# Advanced Technical Depth: DOCX Generation
try:
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

def generate_weekly_digest():
    today = datetime.now()
    filename = f"ClawMentor_Weekly_Report_{today.strftime('%Y-%m-%d')}.docx"
    
    # Matching PPT suggested path
    output_path = Path.home() / ".openclaw" / "reports"
    output_path.mkdir(parents=True, exist_ok=True)
    doc_path = output_path / filename

    if not DOCX_AVAILABLE:
        print("⚠️ python-docx not installed. Skipping .docx generation.")
        return None

    doc = Document()
    doc.add_heading('ClawMentor Bharat - Weekly Progress Report', 0)
    
    # Header
    p = doc.add_paragraph()
    p.add_run(f"Student: Sahana V S | RVCE CSE 4th Year\n").bold = True
    p.add_run(f"Generated: {today.strftime('%d %B %Y')}").italic = True

    doc.add_heading('Summary', level=1)
    doc.add_paragraph("This week showed consistent progress in DSA and project work based on your Living Roadmap.")

    doc.add_heading('Key Achievements', level=1)
    achievements = ["Completed 12 LeetCode problems", "Advanced RAG project structure", "Updated resume with new skills detected by Analyzer"]
    for ach in achievements:
        doc.add_paragraph(ach, style='List Bullet')

    doc.add_heading('Skill Gaps & Next Actions', level=1)
    doc.add_paragraph("• Focus more on System Design\n• Complete Docker deployment module")

    doc.add_heading('Next Week Goals', level=1)
    doc.add_paragraph("1. Build 1 full-stack project\n2. 20+ LeetCode problems\n3. Mock interviews")

    doc.save(doc_path)
    
    print(f"✅ Weekly Report Generated Successfully!")
    print(f"📄 Saved at: {doc_path}")
    
    # Notify user via Telegram/WhatsApp as per PPT
    notify_user("📊 Your Weekly Progress Report is ready! Check the attached DOCX in your reports folder.", "telegram")
    
    return str(doc_path)

if __name__ == "__main__":
    generate_weekly_digest()
